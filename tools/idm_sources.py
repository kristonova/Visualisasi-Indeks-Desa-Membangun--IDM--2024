#!/usr/bin/env python3
r"""Sumber tabular IDM per tahun — dipakai build_idm.py dan audit_join.py.

Setiap tahun punya berkas Excel dengan nama sheet dan kolom yang berbeda.
Modul ini menyeragamkannya menjadi satu bentuk record, dan menerjemahkan kode
wilayah ke penomoran IDM 2024 supaya kedua tahun memakai geometri yang sama.

Satu-satunya perbedaan penomoran antara IDM 2023 dan 2024 adalah Papua Barat
Daya: di IDM 2023 kelima kabupatennya masih berkode 92.xx (Papua Barat),
sedangkan di IDM 2024 sudah 96.xx. Keempat digit desa sesudahnya tidak berubah
(939 dari 939 desa cocok setelah remap).
"""

from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

ROOT = Path(__file__).resolve().parents[1]

NAMA_PROV = {
    "11": "ACEH", "12": "SUMATERA UTARA", "13": "SUMATERA BARAT", "14": "RIAU",
    "15": "JAMBI", "16": "SUMATERA SELATAN", "17": "BENGKULU", "18": "LAMPUNG",
    "19": "KEPULAUAN BANGKA BELITUNG", "21": "KEPULAUAN RIAU", "31": "DKI JAKARTA",
    "32": "JAWA BARAT", "33": "JAWA TENGAH", "34": "DI YOGYAKARTA", "35": "JAWA TIMUR",
    "36": "BANTEN", "51": "BALI", "52": "NUSA TENGGARA BARAT", "53": "NUSA TENGGARA TIMUR",
    "61": "KALIMANTAN BARAT", "62": "KALIMANTAN TENGAH", "63": "KALIMANTAN SELATAN",
    "64": "KALIMANTAN TIMUR", "65": "KALIMANTAN UTARA", "71": "SULAWESI UTARA",
    "72": "SULAWESI TENGAH", "73": "SULAWESI SELATAN", "74": "SULAWESI TENGGARA",
    "75": "GORONTALO", "76": "SULAWESI BARAT", "81": "MALUKU", "82": "MALUKU UTARA",
    "91": "PAPUA", "92": "PAPUA BARAT", "93": "PAPUA SELATAN", "94": "PAPUA TENGAH",
    "95": "PAPUA PEGUNUNGAN", "96": "PAPUA BARAT DAYA",
}

STATUS_TO_CODE = {
    None: 0,
    "": 0,
    "MANDIRI": 1,
    "MAJU": 2,
    "BERKEMBANG": 3,
    "TERTINGGAL": 4,
    "SANGAT TERTINGGAL": 5,
}

# Nilai kolom "pv" di JSON desa. Label harus sama dengan SOURCE_NAME di dashboard.
SOURCE_LABEL = {0: "Reguler", 1: "Update 2023", 2: "Server PDN", 3: "Data 2022"}

# Kab/kota Papua Barat Daya di IDM 2023 (masih di bawah Papua Barat) -> kode IDM 2024.
PBD_2023 = {"9201": "9601", "9204": "9604", "9205": "9605", "9209": "9609", "9210": "9610"}

YEARS: Dict[str, Dict[str, Any]] = {
    "2023": {
        "excel": ROOT / "data" / "indeks-desa-membangun-idm-tahun-2023.xlsx",
        "sheet": "Sheet1",
        "cols": {
            "prov": "KODE_PROV", "prov_name": "PROVINSI",
            "kab": "KODE_KAB_KOT", "kab_name": "KABUPATEN_KOTA",
            "kec": "KODE_KEC", "kec_name": "KECAMATAN",
            "code": "KODE_DESA_DAGRI", "name": "DESA",
            "iks": "IKS_2023", "ike": "IKE_2023", "ikl": "IKL_2023",
            "idm": "NILAI_IDM_2023", "status": "STATUS_IDM_2023", "source": "UPDATE",
        },
        # Kolom UPDATE: tahun pemutakhiran nilai. 2022 = nilai tahun lalu dibawa.
        "sources": {2023: 0, "2023": 0, 2022: 3, "2022": 3},
        "crosswalk": PBD_2023,
    },
    "2024": {
        "excel": ROOT / "data" / "indeks-desa-membangun-tahun-2024-hasil-pemutakhiran.xlsx",
        "sheet": "IDM 2024",
        "cols": {
            "prov": "KODE_PROV", "prov_name": "NAMA_PROVINSI",
            "kab": "KODE_KAB", "kab_name": "NAMA_KABUPATEN",
            "kec": "KODE_KEC", "kec_name": "NAMA_KECAMATAN",
            "code": "KODE_DESA", "name": "NAMA_DESA",
            "iks": "IKS_2024", "ike": "IKE_2024", "ikl": "IKL_2024",
            "idm": "NILAI_IDM_2024", "status": "STATUS_IDM_2024", "source": "Keterangan",
        },
        "sources": {None: 0, "": 0, "Update 2023": 1, "Server PDN": 2},
        "crosswalk": {},
    },
}


def digits(value: Any) -> str:
    """Hanya digit desimal, tanpa lewat float untuk kode panjang."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\D", "", str(value))


def numeric(value: Any) -> Optional[float]:
    """Angka apa adanya (int tetap int) atau None untuk teks/kosong/NaN."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return value if math.isfinite(value) else None


def crosswalk(code: str, year: str) -> str:
    """Kode desa 10 digit sumber tahun `year` -> penomoran IDM 2024."""
    table = YEARS[year]["crosswalk"]
    if len(code) == 10 and code[:4] in table:
        return table[code[:4]] + code[4:]
    return code


def load_rows(year: str, path: Optional[Path] = None) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Baca Excel IDM `year` -> (records, invalid). Kode sudah diterjemahkan."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - hanya pada setup yang rusak
        raise SystemExit(
            "openpyxl belum terpasang. Jalankan dengan .venv\\Scripts\\python "
            "atau pasang openpyxl terlebih dahulu."
        ) from exc

    cfg = YEARS[year]
    path = Path(path or cfg["excel"])
    cols, sources = cfg["cols"], cfg["sources"]
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if cfg["sheet"] not in workbook.sheetnames:
        raise SystemExit(f"Sheet '{cfg['sheet']}' tidak ditemukan di {path}")
    rows = workbook[cfg["sheet"]].iter_rows(values_only=True)
    header = next(rows)
    index = {str(name).strip(): i for i, name in enumerate(header) if name is not None}
    missing = [c for c in cols.values() if c not in index]
    if missing:
        raise SystemExit(f"Kolom Excel {path.name} tidak lengkap: " + ", ".join(missing))

    records: List[Dict[str, Any]] = []
    invalid: List[Dict[str, Any]] = []
    for row_number, row in enumerate(rows, start=2):
        get = lambda key: row[index[cols[key]]] if index[cols[key]] < len(row) else None
        raw = digits(get("code"))
        if not raw and all(v is None for v in row):
            continue  # baris kosong di ekor sheet
        code = crosswalk(raw, year)
        moved = code != raw
        source_value = get("source")
        iks = get("iks")
        record = {
            "code": code,
            "code_source": raw,
            "province_code": code[:2],
            "province": NAMA_PROV.get(code[:2], "") if moved else str(get("prov_name") or "").strip(),
            "regency_code": code[:4],
            "regency": str(get("kab_name") or "").strip(),
            "district_code": code[:6],
            "district": str(get("kec_name") or "").strip(),
            "name": str(get("name") or "").strip(),
            "iks": numeric(iks),
            "ike": numeric(get("ike")),
            "ikl": numeric(get("ikl")),
            "idm": numeric(get("idm")),
            "status": STATUS_TO_CODE.get(get("status"), -1),
            "source": sources.get(source_value, -1),
            "source_label": SOURCE_LABEL.get(sources.get(source_value, -1), str(source_value)),
            # Sel skor berisi teks (mis. alasan tidak dinilai) -> simpan sebagai catatan.
            "source_note": str(iks).strip() if iks not in (None, "") and numeric(iks) is None else None,
            "row": row_number,
        }
        # Kolom kode induk harus konsisten dengan awalan kode desa sumber.
        prefix_ok = (
            digits(get("prov")).zfill(2) == raw[:2]
            and digits(get("kab")).zfill(4) == raw[:4]
            and digits(get("kec")).zfill(6) == raw[:6]
        )
        if len(code) != 10 or not code.isdigit() or not prefix_ok:
            invalid.append({"row": row_number, "code": raw, "name": record["name"],
                            "reason": "kode bukan 10 digit" if len(code) != 10 else "kode induk tidak konsisten"})
        records.append(record)
    workbook.close()
    return records, invalid
